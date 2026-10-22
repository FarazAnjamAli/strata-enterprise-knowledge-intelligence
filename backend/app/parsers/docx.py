from pathlib import Path

from docx import Document

from app.parsers.base import ParsedBlock, ParseError


def parse(path: Path) -> list[ParsedBlock]:
    try:
        doc = Document(str(path))
    except Exception as exc:
        raise ParseError(f"DOCX parse failed: {exc}") from exc

    blocks: list[ParsedBlock] = []
    section = None
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if para.style.name.startswith("Heading"):
            section = text
            continue
        blocks.append(ParsedBlock(text=text, section=section))

    for t_idx, table in enumerate(doc.tables):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        text = "\n".join(r for r in rows if r.strip(" |"))
        if text:
            blocks.append(ParsedBlock(text=text, section=f"table {t_idx + 1}", metadata={"table": t_idx + 1}))
    return blocks
